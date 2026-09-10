#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成 result2.xlsx（流式写入）。

结果矩阵为 2 张工作表 x 219685 行 x 22 列，合计约 970 万个单元格。
spreadsheets 技能推荐的 @oai/artifact-tool 在该规模下内存耗尽（构建 15 分钟后
崩溃），因此这里改用 openpyxl 的 write_only 流式模式 —— 它专为大表设计，
峰值内存与行数无关。其余几个 result 文件规模较小，仍由 artifact-tool 生成。

输入: 输出/result2.bin、输出/result_meta.json
输出: 输出/result2.xlsx
"""

from __future__ import annotations

import json
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

PROJECT = Path(__file__).resolve().parent.parent
OUT = PROJECT / "输出"

HEADER_LABEL = "时间\\到药材中心的距离"
NUMFMT = "0.0000"


def write_sheet(sheet, times, grid, positions):
    sheet.append([HEADER_LABEL, *positions])
    n_pos = len(positions)
    for i in range(len(times)):
        base = i * n_pos
        row = [float(times[i])]
        row.extend(np.round(grid[base : base + n_pos], 4).tolist())
        sheet.append(row)


def style(sheet, n_cols):
    sheet.column_dimensions["A"].width = 23
    for c in range(2, n_cols + 1):
        sheet.column_dimensions[get_column_letter(c)].width = 10
        sheet.column_dimensions[get_column_letter(c)].number_format = NUMFMT
    sheet.freeze_panes = "B2"
    sheet.sheet_view.showGridLines = False


def main():
    meta = json.loads((OUT / "result_meta.json").read_text(encoding="utf-8"))
    positions = meta["positions_cm"]
    n = meta["n2"]
    n_pos = len(positions)

    arr = np.fromfile(OUT / "result2.bin", dtype="<f8")
    times = arr[:n]
    temp = arr[n : n + n * n_pos]
    conc = arr[n + n * n_pos : n + 2 * n * n_pos]

    wb = Workbook(write_only=True)
    for name, grid in (("温度", temp), ("水分浓度", conc)):
        ws = wb.create_sheet(name)
        write_sheet(ws, times, grid, positions)
        style(ws, n_pos + 1)

    target = OUT / "result2.xlsx"
    wb.save(target)
    print("已生成: %s  (%d 行 x %d 列 x 2 表)" % (target, n, n_pos + 1))


if __name__ == "__main__":
    main()
